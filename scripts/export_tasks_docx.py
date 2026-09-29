#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export uploaded quests (q_doc_*, q_cmp_*) to Word for editing."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
QUESTS_PATH = ROOT / "assets" / "content" / "quests.json"
OUT = ROOT / "Задания_для_редактирования.docx"

THEME_RU = {
    "budget": "бюджет",
    "savings": "накопления",
    "purchases": "покупки",
}
DIFF_RU = {
    "easy": "лёгкий",
    "medium": "средний",
    "hard": "сложный",
}
AGE_RU = {
    "junior": "младшие",
    "middle": "средние",
    "senior": "старшие",
}


def set_run_font(run, name="Times New Roman", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_para(doc, text, bold=False, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_meta(doc, quest):
    theme = THEME_RU.get(quest.get("theme", ""), quest.get("theme", ""))
    diff = DIFF_RU.get(quest.get("difficulty", ""), quest.get("difficulty", ""))
    ages = ", ".join(AGE_RU.get(a, a) for a in quest.get("ageTracks", []))
    add_para(
        doc,
        f"ID: {quest['id']}  |  тема: {theme}  |  сложность: {diff}  |  возраст: {ages or 'все'}",
        size=10,
    )


def add_quest(doc, index, quest):
    add_heading(doc, f"{index}. {quest.get('title', quest['id'])}", level=2)
    add_meta(doc, quest)
    add_para(doc, "Вопрос / условие:", bold=True)
    add_para(doc, quest.get("prompt", ""))
    tip = quest.get("tip")
    if tip:
        add_para(doc, "Подсказка:", bold=True)
        add_para(doc, tip)

    interaction = quest.get("interaction")
    if interaction and interaction != "choice":
        add_para(doc, f"Тип взаимодействия: {interaction}", bold=True)

    sort_items = quest.get("sortItems") or []
    if sort_items:
        add_para(doc, "Карточки для сортировки:", bold=True)
        for item in sort_items:
            add_para(
                doc,
                f"  • {item.get('label', item.get('id', ''))} → корзина «{item.get('correctBucket', '')}»",
            )

    options = quest.get("options") or []
    if options:
        add_para(doc, "Варианты ответа:", bold=True)
        for opt in options:
            mark = "✓ верный" if opt.get("isGood") else "✗ неверный"
            reward = opt.get("reward", "")
            label = opt.get("label", "")
            add_para(doc, f"  [{mark}] {label}  (награда: {reward})")
            expl = opt.get("explanation")
            if expl:
                add_para(doc, f"      Пояснение: {expl}", size=11)

    # Blank lines for editor notes
    add_para(doc, "— правки / комментарии: ________________________________", size=10)
    doc.add_paragraph()


def main():
    quests = json.loads(QUESTS_PATH.read_text(encoding="utf-8"))
    uploaded = [
        q
        for q in quests
        if q.get("id", "").startswith("q_doc_") or q.get("id", "").startswith("q_cmp_")
    ]
    doc_part = [q for q in uploaded if q["id"].startswith("q_doc_")]
    cmp_part = [q for q in uploaded if q["id"].startswith("q_cmp_")]

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(
        "Задания для редактирования\n"
        "«Питомец Финни» — контент, загруженный из документов"
    )
    set_run_font(r, size=16, bold=True)

    add_para(
        doc,
        f"Всего заданий: {len(uploaded)} "
        f"(из документа: {len(doc_part)}, сравнение чисел: {len(cmp_part)}). "
        "Источник в приложении: assets/content/quests.json. "
        "Отредактируйте тексты здесь, затем внесите правки в JSON.",
        size=11,
    )
    doc.add_paragraph()

    # Summary table
    add_heading(doc, "Сводка", level=1)
    table = doc.add_table(rows=1 + len(uploaded), cols=5)
    table.style = "Table Grid"
    headers = ["№", "ID", "Название", "Тема", "Сложность"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=9, bold=True)
    for i, q in enumerate(uploaded, start=1):
        row = table.rows[i].cells
        vals = [
            str(i),
            q.get("id", ""),
            q.get("title", ""),
            THEME_RU.get(q.get("theme", ""), q.get("theme", "")),
            DIFF_RU.get(q.get("difficulty", ""), q.get("difficulty", "")),
        ]
        for c_i, val in enumerate(vals):
            row[c_i].text = val
            for p in row[c_i].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()

    add_heading(doc, f"Часть 1. Задания из документа ({len(doc_part)})", level=1)
    for i, q in enumerate(doc_part, start=1):
        add_quest(doc, i, q)

    add_heading(doc, f"Часть 2. Сравнение чисел ({len(cmp_part)})", level=1)
    for i, q in enumerate(cmp_part, start=1):
        add_quest(doc, i, q)

    doc.save(OUT)
    print(f"Saved: {OUT}")
    print(f"Quests: {len(uploaded)} (doc={len(doc_part)}, cmp={len(cmp_part)})")


if __name__ == "__main__":
    main()
